"""Genera los tres manuales editables de Re-Fluye a partir del Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCES = ROOT / "refluye" / "manuales"
OUTPUT = ROOT / "output" / "word"
BLUE = RGBColor(18, 72, 102)
TEAL = RGBColor(0, 112, 116)
GRAY = RGBColor(86, 98, 106)


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def margins(cell, top=85, start=110, bottom=85, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, amount in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        elt = OxmlElement(f"w:{side}")
        elt.set(qn("w:w"), str(amount))
        elt.set(qn("w:type"), "dxa")
        mar.append(elt)
    tc_pr.append(mar)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    relation = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "007074")
    props.append(color)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = label
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def inline(paragraph, value: str) -> None:
    for part in INLINE.split(value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = part[1:-1].split("](", 1)
            if url.startswith("http"):
                add_hyperlink(paragraph, label, url)
            else:
                paragraph.add_run(label)
        else:
            paragraph.add_run(part)


def base_doc(visual: bool = False) -> Document:
    doc = Document()
    section = doc.sections[0]
    if visual:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.25)
        section.bottom_margin = Cm(1.1)
        section.left_margin = Cm(1.35)
        section.right_margin = Cm(1.35)
    else:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9 if visual else 10.5)
    normal.font.color.rgb = RGBColor(25, 42, 51)
    normal.paragraph_format.space_after = Pt(4 if visual else 7)
    normal.paragraph_format.line_spacing = 1.08 if visual else 1.15
    for style, size, color, before, after in (
        ("Title", 22 if visual else 26, RGBColor(0, 0, 0), 0, 10),
        ("Heading 1", 12 if visual else 16, BLUE, 11, 5),
        ("Heading 2", 10 if visual else 12, TEAL, 8, 3),
        ("Heading 3", 9 if visual else 10.5, BLUE, 6, 2),
    ):
        s = styles[style]
        s.font.name = "Aptos Display" if style != "Heading 3" else "Aptos"
        s.font.size = Pt(size)
        s.font.bold = style != "Title"
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("RE-FLUYE  ·  02 OCT 2026").font.color.rgb = GRAY
    return doc


def add_table(doc: Document, rows: list[list[str]], visual: bool) -> None:
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.autofit = True
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for idx, row in enumerate(rows):
        new_row = table.add_row()
        cells = new_row.cells
        cant_split = OxmlElement("w:cantSplit")
        new_row._tr.get_or_add_trPr().append(cant_split)
        if idx == 0:
            repeat = OxmlElement("w:tblHeader")
            new_row._tr.get_or_add_trPr().append(repeat)
        for cell, value in zip(cells, row):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell, 55 if visual else 85, 70 if visual else 110, 55 if visual else 85, 70 if visual else 110)
            if idx == 0:
                shade(cell, "DCEDEF")
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            inline(para, value)
            for run in para.runs:
                run.font.size = Pt(7.2 if visual else 8.5)
                if idx == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_markdown(doc: Document, path: Path, *, visual: bool = False) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if line.startswith("```"):
            if in_code:
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(0.35)
                para.paragraph_format.space_after = Pt(6)
                run = para.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                code_lines.clear()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(lines[i - 1])
            continue
        if not line or (visual and (line.startswith("**Formato de entrega:") or line.startswith("**Edición de prototipo"))):
            continue
        if line.startswith("|"):
            rows = []
            while True:
                if re.fullmatch(r"\|[\s:|\-]+\|", line):
                    pass
                else:
                    rows.append([v.strip() for v in line.strip("|").split("|")])
                if i >= len(lines) or not lines[i].strip().startswith("|"):
                    break
                line = lines[i].strip()
                i += 1
            if rows:
                add_table(doc, rows, visual)
            continue
        if line.startswith("# "):
            inline(doc.add_paragraph(style="Title"), re.sub(r"[·]", "", line[2:]).strip())
        elif line.startswith("## "):
            title = line[3:].replace(" · ", "  ")
            if visual and title.startswith("Cara 2"):
                doc.add_page_break()
            inline(doc.add_paragraph(style="Heading 1"), title)
        elif line.startswith("### "):
            inline(doc.add_paragraph(style="Heading 2"), line[4:].replace(" · ", "  "))
        elif re.match(r"^\d+\. ", line):
            para = doc.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(2 if visual else 5)
            inline(para, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            para = doc.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(2 if visual else 5)
            inline(para, line[2:])
        else:
            inline(doc.add_paragraph(), line)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for stem, output_name, visual in (
        ("01-guia-app", "refluye-guia-app.docx", False),
        ("02-guia-visual-equipo", "refluye-guia-visual-equipo-a4.docx", True),
        ("03-manual-tecnico", "refluye-manual-tecnico.docx", False),
    ):
        doc = base_doc(visual)
        parse_markdown(doc, SOURCES / f"{stem}.md", visual=visual)
        target = OUTPUT / output_name
        doc.save(target)
        print(target)


if __name__ == "__main__":
    main()
